"""Regenerate the binary test fixtures: python3 test/fixtures/make_fixtures.py
Needs python-docx and python-pptx (pip install python-docx python-pptx)."""
from pathlib import Path

import docx
import pptx

HERE = Path(__file__).parent


def make_pdf(path: Path) -> None:
    """Three pages: two with text, one with no text layer."""
    pages = [b"BT /F1 18 Tf 72 720 Td (Projections and transformations) Tj 0 -30 Td (Pinhole camera model) Tj ET",
             b"BT /F1 14 Tf 72 720 Td (Homogeneous coordinates) Tj ET",
             b""]
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % (4 + 2 * i) for i in range(len(pages))), len(pages)),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, content in enumerate(pages):
        objs.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (5 + 2 * i))
        objs.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content))
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (n, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(bytes(out))


def make_docx(path: Path) -> None:
    d = docx.Document()
    d.add_heading("Lab 3 answers", 1)
    d.add_paragraph("Name: ____ & co <x>")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Q1", "answer"
    t.cell(1, 0).text, t.cell(1, 1).text = "Q2", "b"
    d.add_heading("Part 2", 1)
    d.add_paragraph("Second section text")
    d.save(path)


def make_pptx(path: Path) -> None:
    p = pptx.Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Projections"
    s.placeholders[1].text = "Pinhole camera\nHomogeneous coords"
    s.notes_slide.notes_text_frame.text = "Say something about focal length"
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Second"
    p.save(path)


if __name__ == "__main__":
    make_pdf(HERE / "sample.pdf")
    make_docx(HERE / "sample.docx")
    make_pptx(HERE / "sample.pptx")
